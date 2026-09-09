import os
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import unittest
import openpyxl
from datetime import date
import tkinter as tk

from gui_ocm_net import OCMNetGUI
from generate_ocm_net_report import (
    process_vendor_files,
    read_vendor_file,
    FILE_CONFIGS,
    KPI_MAP,
)


class TestGUIOCMNet(unittest.TestCase):
    def test_main_vendor_labels_contains_nokia_2g(self):
        """Verify ('N_2G', 'Nokia 2G') is in MAIN_VENDOR_LABELS."""
        keys = [k for k, _ in OCMNetGUI.MAIN_VENDOR_LABELS]
        self.assertIn('N_2G', keys)
        self.assertEqual(
            dict(OCMNetGUI.MAIN_VENDOR_LABELS)['N_2G'],
            'Nokia 2G'
        )

    def test_gui_initialization(self):
        """Verify OCMNetGUI initializes without errors in headless mode."""
        root = tk.Tk()
        root.withdraw()
        try:
            gui = OCMNetGUI(root)
            self.assertIn('N_2G', gui.file_vars)
            self.assertIsInstance(gui.file_vars['N_2G'], tk.StringVar)
        finally:
            root.destroy()

    def test_auto_detect_nokia_2g(self):
        """Test auto-detection logic for Nokia 2G raw file."""
        import tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            n2g_file = os.path.join(tmpdir, 'N 2G_20260510.xlsx')
            open(n2g_file, 'w').close()

            root = tk.Tk()
            root.withdraw()
            try:
                gui = OCMNetGUI(root)
                gui.raw_dir_var.set(tmpdir)
                gui._auto_detect()
                self.assertEqual(gui.file_vars['N_2G'].get(), n2g_file)
            finally:
                root.destroy()

    def test_process_vendor_files_nokia_2g(self):
        """Test reading a dummy Nokia 2G Excel file and updating report."""
        import tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            # Create a mock Nokia 2G vendor excel file
            n2g_path = os.path.join(tmpdir, 'N 2G_sample.xlsx')
            wb_n2g = openpyxl.Workbook()
            ws_n2g = wb_n2g.active

            # Nokia 2G format: Row 1 = Headers, Row 2 = KPI Codes, Row 3+ = Data
            ws_n2g.append([
                'Period start time',
                'BCF name',
                'TCH Availability Normal TRXs',
                'Erlang_Traffic_Carried_2G',
                'ORA_2G_CSSR_CS_new',
                'ORA_2G_Call_Drop_CS_new',
            ])
            ws_n2g.append([
                'PSTART', 'BCF', 'K1', 'K2', 'K3', 'K4'
            ])
            ws_n2g.append([
                '2026-05-10 00:00:00',
                'ADM_004_N_SITE',
                99.5,
                1500.0,  # 1500 Erl -> 1.5 Kerl
                98.2,
                0.8,
            ])
            wb_n2g.save(n2g_path)

            # Test read_vendor_file directly
            vdata = read_vendor_file(n2g_path, 'N_2G')
            self.assertIn('ADM_004', vdata)
            d = date(2026, 5, 10)
            self.assertIn(d, vdata['ADM_004'])
            site_kpis = vdata['ADM_004'][d]
            self.assertEqual(site_kpis.get('Avail2G'), 99.5)
            self.assertEqual(site_kpis.get('DailyCombinedCSTrafic (Kerl)'), 1.5)
            self.assertEqual(site_kpis.get('CSSR2G'), 98.2)
            self.assertEqual(site_kpis.get('DCR2G'), 0.8)

            # Create mock OCM Daily template file
            ocm_daily_path = os.path.join(tmpdir, 'OCM_Daily.xlsx')
            wb_ocm = openpyxl.Workbook()
            ws_avail = wb_ocm.active
            ws_avail.title = 'Avail2G'
            ws_avail.append(['Title Row'])
            ws_avail.append(['ColA', 'Code du Site', 'ColC', 'ColD', 'ColE', 'ColF', 'ColG', 'ColH', date(2026, 5, 10)])
            ws_avail.append(['Name', 'ADM_004', '', '', '', '', '', '', None])

            # Add other sheets
            for sheet_name in ['DailyCombinedCSTrafic (Kerl)', 'CSSR2G', 'DCR2G']:
                ws = wb_ocm.create_sheet(title=sheet_name)
                ws.append(['Title Row'])
                ws.append(['ColA', 'Code du Site', 'ColC', 'ColD', 'ColE', 'ColF', 'ColG', 'ColH', date(2026, 5, 10)])
                ws.append(['Name', 'ADM_004', '', '', '', '', '', '', None])

            wb_ocm.save(ocm_daily_path)

            # Process vendor files
            process_vendor_files(
                vendor_files={'N_2G': n2g_path},
                ocm_files={'daily': ocm_daily_path},
                granularities={'daily'},
            )

            # Check that values were written into the OCM Daily file
            wb_res = openpyxl.load_workbook(ocm_daily_path, data_only=True)
            self.assertEqual(wb_res['Avail2G'].cell(row=3, column=9).value, 99.5)
            self.assertEqual(wb_res['DailyCombinedCSTrafic (Kerl)'].cell(row=3, column=9).value, 1.5)
            self.assertEqual(wb_res['CSSR2G'].cell(row=3, column=9).value, 98.2)
            self.assertEqual(wb_res['DCR2G'].cell(row=3, column=9).value, 0.8)


if __name__ == '__main__':
    unittest.main()
